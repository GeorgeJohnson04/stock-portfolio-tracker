from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties

wb = load_workbook('Finance Project.xlsx')
ws_dash = wb['Dashboard']

ws_dash._charts.clear()

type_colors = ['1F4E79', '2E75B6', '00B0A0', '7B8FA1']
ind_colors = ['1F4E79', '2E75B6', '00B0A0', 'D4A017', '2ECC71', 'E74C3C', '8E44AD', '7B8FA1']

def apply_slice_colors(series, colors):
    for i, color in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = color
        series.data_points.append(pt)

def make_label_props():
    lbl = DataLabelList()
    lbl.showPercent = True
    lbl.showCatName = True
    lbl.showVal = False
    lbl.showSerName = False
    lbl.showLeaderLines = True
    lbl.numFmt = '0%'
    lbl.dLblPos = 'outEnd'
    lbl.separator = '\n'
    fp = CharacterProperties(sz=1000, b=True)
    fp.solidFill = '333333'
    pp = ParagraphProperties(defRPr=fp)
    lbl.textProperties = RichText(p=[Paragraph(pPr=pp, endParaRPr=fp)])
    return lbl

# ============================================================
# PIE 1: By Position Type
# ============================================================
pie1 = PieChart()
pie1.title = 'By Position Type'
pie1.style = 2
pie1.width = 18
pie1.height = 13

data1 = Reference(ws_dash, min_col=2, min_row=14, max_row=18)
cats1 = Reference(ws_dash, min_col=1, min_row=15, max_row=18)
pie1.add_data(data1, titles_from_data=True)
pie1.set_categories(cats1)
pie1.dataLabels = make_label_props()
apply_slice_colors(pie1.series[0], type_colors)
pie1.legend.position = 'b'

ws_dash.add_chart(pie1, 'E3')

# ============================================================
# PIE 2: By Industry
# ============================================================
pie2 = PieChart()
pie2.title = 'By Industry'
pie2.style = 2
pie2.width = 18
pie2.height = 13

data2 = Reference(ws_dash, min_col=2, min_row=29, max_row=37)
cats2 = Reference(ws_dash, min_col=1, min_row=30, max_row=37)
pie2.add_data(data2, titles_from_data=True)
pie2.set_categories(cats2)
pie2.dataLabels = make_label_props()
apply_slice_colors(pie2.series[0], ind_colors)
pie2.legend.position = 'b'

ws_dash.add_chart(pie2, 'E18')

wb.save('Finance Project.xlsx')
print('Done - 2 clean pie charts saved')
